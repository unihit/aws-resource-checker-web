import ast
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys
import unittest
from tempfile import TemporaryDirectory
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT / 'collector'))
import checker
import war

REGION = 'ap-northeast-2'


def snapshot(service='EBS', data=None, status='OK', region=REGION):
    row = checker.resource('test','000000000000',region,service,'test-resource',data or {})
    report = dict(version='0.1.0',mode='DEMO',started_at='2026-10-23T00:00:00Z',status='COMPLETE',
                  requested_regions=[REGION],requested_services=[service],resources=[row], findings=[],
                  coverage=[dict(profile='test',account='000000000000',region=region,service=service,
                    operation='demo',status=status,count=1,error='' if status=='OK' else 'AccessDenied')])
    return report


def result(report, rule, identifier='test-resource'):
    return next(f for f in war.evaluate(report)['findings'] if f['rule']==rule and f['id']==identifier)


class Rules(unittest.TestCase):
    def test_boolean_and_number_missing_types(self):
        for v in (None,'true','false',0,1,[],{}):
            self.assertIsNone(war.boolean(v))
        self.assertIsNone(war.number(True,lambda n:n==1))
        self.assertTrue(war.boolean(True))

    def test_every_rule_missing_fields_is_unknown(self):
        for rule,_,service,*_ in war.RULES:
            with self.subTest(rule=rule):
                r=snapshot(service)
                if service=='S3':
                    r['coverage'][0]['operation']='bucket_detail:test-resource'
                self.assertEqual(result(r,rule)['status'],'UNKNOWN')

    def test_no_network_and_no_mutation(self):
        r=checker.demo(); original=deepcopy(r)
        with patch('boto3.Session',side_effect=AssertionError('No AWS')):
            value=war.evaluate(r)
        self.assertEqual(r,original)
        self.assertEqual({f['pillar'] for f in value['findings']},set(war.PILLARS))
        self.assertEqual(sum(f['status']=='MANUAL' for f in value['findings']),6)

    def test_encrypted_pass_fail(self):
        self.assertEqual(result(snapshot(data={'Encrypted':True}),'SEC-001')['status'],'PASS')
        self.assertEqual(result(snapshot(data={'Encrypted':False}),'SEC-001')['status'],'FAIL')

    def test_failed_or_missing_coverage_never_pass(self):
        for state in ('ERROR','PARTIAL','LIMIT','UNKNOWN'):
            with self.subTest(state=state):
                self.assertEqual(result(snapshot(data={'Encrypted':True},status=state),'SEC-001')['status'],'UNKNOWN')
        r=snapshot(data={'Encrypted':True});r['coverage']=[]
        self.assertEqual(result(r,'SEC-001')['status'],'UNKNOWN')

    def test_partial_known_fail_and_unknown_scope(self):
        r=snapshot(data={'Encrypted':False},status='PARTIAL')
        value=war.evaluate(r)
        self.assertEqual(result(r,'SEC-001')['status'],'FAIL')
        self.assertTrue(any(f['rule']=='DATA-001' and f['status']=='UNKNOWN' for f in value['findings']))

    def test_scope_empty_vs_failure(self):
        r=snapshot();r['resources']=[];r['coverage'][0]['count']=0
        self.assertEqual(result(r,'SEC-001','scope')['status'],'NOT_APPLICABLE')
        r['coverage'][0]['status']='ERROR'
        self.assertEqual(result(r,'SEC-001','scope')['status'],'UNKNOWN')

    def test_regional_missing_scope(self):
        r=snapshot(data={'Encrypted':True});r['requested_regions'].append('us-east-1')
        rows=war.evaluate(r)['findings']
        self.assertTrue(any(f['rule']=='SEC-001' and f['region']=='us-east-1' and f['status']=='UNKNOWN' for f in rows))

    def test_scope_missing_nonempty_population(self):
        r=snapshot();r['resources']=[]
        self.assertEqual(result(r,'SEC-001','scope')['status'],'UNKNOWN')

    def test_basic_flags_never_replace_missing_or_conflicting_fields(self):
        r=snapshot();r['findings']=[{k:r['resources'][0][k] for k in ('profile','account','region','service','id')}|{'rule':'EBS_ENCRYPTION'}]
        self.assertEqual(result(r,'SEC-001')['status'],'UNKNOWN')
        r['resources'][0]['details']['Encrypted']=True
        self.assertEqual(result(r,'SEC-001')['status'],'UNKNOWN')
        r['resources'][0]['details']['Encrypted']=False
        self.assertEqual(result(r,'SEC-001')['status'],'FAIL')

    def test_security_groups_and_ipv6(self):
        for cidr in ('0.0.0.0/0','::/0'):
            group,key=('Ipv6Ranges','CidrIpv6') if ':' in cidr else ('IpRanges','CidrIp')
            p={'IpProtocol':'tcp','FromPort':1,'ToPort':65535,group:[{key:cidr}]}
            self.assertFalse(war.admin_closed({'IpPermissions':[p]}))
            p['FromPort']=443;p['ToPort']=443
            self.assertTrue(war.admin_closed({'IpPermissions':[p]}))
            p['IpProtocol']='-1'
            self.assertFalse(war.admin_closed({'IpPermissions':[p]}))
        self.assertIsNone(war.admin_closed({'IpPermissions':[{'IpProtocol':'tcp','IpRanges':[{'CidrIp':'bad'}]}]}))
        self.assertIsNone(war.admin_closed({'IpPermissions':[{'IpProtocol':'tcp','IpRanges':[{'CidrIp':'0.0.0.0/0'}]}]}))

    def test_rds_threshold_aurora_and_az(self):
        for days,status in ((0,'FAIL'),(6,'FAIL'),(7,'PASS'),(8,'PASS'),('7','UNKNOWN')):
            self.assertEqual(result(snapshot('RDS',{'BackupRetentionPeriod':days}),'REL-002')['status'],status)
        self.assertIsNone(war.multi_az({'Engine':'aurora-postgresql','MultiAZ':True}))
        self.assertTrue(war.azs({'AvailabilityZones':[{'ZoneName':'a'},{'ZoneName':'b'}]}))
        self.assertFalse(war.azs({'AvailabilityZones':[{'ZoneName':'a'},{'ZoneName':'a'}]}))
        self.assertIsNone(war.azs({'AvailabilityZones':[{}]}))

    def test_monitoring_inactive_na(self):
        self.assertEqual(war.monitoring({'State':{'Name':'stopped'}}),'NOT_APPLICABLE')
        self.assertIsNone(war.monitoring({'State':{'Name':'running'}}))
        self.assertTrue(war.monitoring({'State':{'Name':'running'},'Monitoring':{'State':'enabled'}}))

    def test_bpa_missing_bucket_detail_even_if_list_ok(self):
        data={'PublicAccessBlock':{k:True for k in ('BlockPublicAcls','IgnorePublicAcls','BlockPublicPolicy','RestrictPublicBuckets')}}
        r=snapshot('S3',data)
        r['coverage'][0]['operation']='list_buckets';r['coverage'][0]['region']='global'
        self.assertEqual(result(r,'SEC-005')['status'],'UNKNOWN')
        detail=deepcopy(r['coverage'][0]);detail.update(operation='bucket_detail:test-resource',status='ERROR',error='AccessDenied')
        r['coverage'].append(detail)
        self.assertEqual(result(r,'SEC-005')['status'],'UNKNOWN')
        r['coverage'][-1]['status']='OK'
        self.assertEqual(result(r,'SEC-005')['status'],'PASS')

    def test_uncollected_services_unknown(self):
        rows=war.evaluate(checker.demo())['findings']
        for rule,*_ in war.UNAVAILABLE:
            self.assertTrue(all(f['status']=='UNKNOWN' for f in rows if f['rule']==rule))

    def test_evidence_does_not_copy_tags_or_raw_payload(self):
        r=snapshot(data={'Encrypted':True,'Tags':[{'Key':'Secret','Value':'DO-NOT-COPY'}]})
        self.assertNotIn('DO-NOT-COPY',json.dumps(war.evaluate(r)))
        r=snapshot('SecurityGroup',{'IpPermissions':[{'IpProtocol':'tcp','FromPort':22,'ToPort':22,
            'IpRanges':[{'CidrIp':'0.0.0.0/0','Description':'DO-NOT-COPY'}]}]})
        self.assertNotIn('DO-NOT-COPY',json.dumps(war.evaluate(r)))


class Integration(unittest.TestCase):
    def test_auth_collection_and_legacy_rules_unchanged(self):
        # AST fingerprints from the pinned upstream base 745eec0; formatting is ignored.
        expected={
            'checker.py':{'resource':'d4218ad0eafb9fc52b57698cff2448107843ad485481ab78572119f4231622b7',
                          'evaluate':'ea9e3636eed7270c3dd55ec18c40b7adde9b8ee6837ae041a2e4f61500fb1b8d',
                          'scan':'238506bd61b048fc4b983e3d484a99e1eb28da239c31c73199c73be26982b755',
                          'inventory_fields':'2e0fb00cac04601ae5508b874db7e71454e1dfeaba4240c99c0d2bc9f8aa1f9e'},
            'cloudshell.py':{'main':'a6be9b1b7a1fd6059c7df0e6a201f99d37581c82e4920a1730a46a3aea5a788f',
                             'parse_regions':'8a4d6f19a72545ec8dd1af18bd56e916850bd749233d0b56602185f70f99f90a'}}
        for filename,functions in expected.items():
            tree=ast.parse((ROOT/'collector'/filename).read_text(encoding='utf-8'))
            def canonical(node):
                if isinstance(node,ast.AST):
                    return [type(node).__name__,{k:canonical(v) for k,v in ast.iter_fields(node) if v is not None and v != []}]
                if isinstance(node,list):
                    return [canonical(v) for v in node]
                return node
            actual={n.name:hashlib.sha256(json.dumps(canonical(n),sort_keys=True).encode()).hexdigest() for n in tree.body if isinstance(n,ast.FunctionDef)}
            for name,digest in functions.items():
                self.assertEqual(actual[name],digest,filename+':'+name)

    def test_demo_html_excel_json_and_legacy_sheets(self):
        from openpyxl import load_workbook
        r=checker.demo();original=deepcopy(r)
        with TemporaryDirectory() as directory:
            folder=checker.export_report(r,directory)
            saved=json.loads((folder/'report.json').read_text(encoding='utf-8'))
            self.assertEqual(r,original)
            for key in ('resources','findings','coverage','status'):
                self.assertEqual(saved[key],original[key])
            self.assertIn('war',saved)
            workbook=load_workbook(folder/'report.xlsx')
            self.assertEqual(workbook.sheetnames[:8],['Summary','EBS','RDS','SecurityGroup','EKS','Findings','Coverage','Inventory'])
            self.assertEqual(workbook.sheetnames[-3:],['WAR Summary','WAR Findings','WAR Manual'])
            self.assertEqual(workbook['EBS']['C8'].value,'vol-demo')
            self.assertEqual(workbook['EBS']['F8'].value,'available')
            self.assertEqual(len(workbook['Summary']._charts),1)
            self.assertEqual(workbook['WAR Findings'].max_row,7+len(saved['war']['findings']))
            self.assertIn('WAR 사전 진단',(folder/'report.html').read_text(encoding='utf-8'))

    def test_excel_html_literal_text(self):
        from openpyxl import load_workbook
        r=snapshot(data={'Encrypted':False})
        r['resources'][0]['id']='=HYPERLINK("https://bad")'
        r['resources'][0]['details']['State']='available'
        r['resources'][0]['details']['Attachments']=[]
        r['resources'][0]['name']='<script>alert(1)</script>'
        with TemporaryDirectory() as directory:
            folder=checker.export_report(r,directory)
            workbook=load_workbook(folder/'report.xlsx')
            self.assertEqual(workbook['EBS']['C8'].data_type,'s')
            self.assertEqual(workbook['WAR Findings']['H8'].data_type,'s')
            self.assertNotIn('<script>alert(1)</script>',(folder/'report.html').read_text(encoding='utf-8'))

    def test_manifest_matches_actual_three_files_and_no_new_permissions(self):
        page=(ROOT/'index.html').read_text(encoding='utf-8')
        data=json.loads(re.search(r'window\.CHECKER_DATA = (\{[^\n]+\});',page).group(1))
        self.assertEqual(set(data['collector']),{'checker.py','cloudshell.py','war.py'})
        for filename,digest in data['collector'].items():
            self.assertEqual(digest,hashlib.sha256((ROOT/'collector'/filename).read_bytes()).hexdigest())
        self.assertFalse(any('cloudtrail:' in action.lower() or 'cloudwatch:' in action.lower() for stmt in data['permissions']['Statement'] for action in stmt['Action']))
        self.assertEqual(data['sample']['war']['counts'],war.evaluate(data['sample'])['counts'])

    def test_offline_cli_preserves_source(self):
        r=checker.demo()
        with TemporaryDirectory() as directory:
            before=Path(directory)/'before.json';after=Path(directory)/'after.json'
            before.write_text(json.dumps(r),encoding='utf-8')
            with patch.object(sys,'argv',['war.py',str(before),'--out',str(after)]):
                war.main()
            self.assertEqual(json.loads(before.read_text(encoding='utf-8')),r)
            self.assertIn('war',json.loads(after.read_text(encoding='utf-8')))


if __name__=='__main__':
    unittest.main()
